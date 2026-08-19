#!/usr/bin/env python3
"""ingest.py — 문서 디렉터리 → 장착 가능한 카트리지 (지식 YAML) 배치 인제스터.

한 번 실행하면 디렉터리 안 파일들을 포맷별로 텍스트 추출하고,
  · 이미 Q&A 구조면(csv/json/xlsx의 question·answer 열) → 결정적 직접 매핑(무손실)
  · 비정형이면 → 설치된 로컬 LLM(/api/wizard/knowledge-convert)이 Q&A 초안 생성
그 결과를 cartridges/<이름>/knowledge/*.yaml + cartridge.yaml 로 쓰고 validate 한다.

새 변환 메커니즘을 만들지 않는다 — 위저드의 /knowledge-convert 와
cartridge_validate 를 배치로 감싸는 얇은 층일 뿐이다.

⚠️ 두 가지 정직한 경계:
  1) 텍스트 계열(csv/json/xml/md/txt)은 바로. PDF/DOCX/XLSX/이미지는 추출 라이브러리가
     설치돼 있어야 하며(아래 안내), 미설치면 그 파일만 건너뛴다.
  2) 비정형 문서의 LLM 변환 결과는 **초안**이다 — validate 가 stub(빈 항목)은 거르지만
     "답이 맞는지"는 검증하지 못한다. 장착 전 검수 권장.

사용:
  python ingest.py <소스_디렉터리> <카트리지_이름> [옵션]
    --count N            비정형 문서 청크당 뽑을 Q&A 개수 (기본 6)
    --base-url URL       콘솔 주소 (기본 config.ini [server] 포트로 https://localhost)
    --key KEY            Bearer 키 (기본 api_keys/default.key)
    --system-prompt-file F  프롬프트 슬롯에 넣을 system.txt (없으면 기본 문구)
    --dry-run            추출·변환만 하고 파일 쓰지 않음 (개수만 보고)
"""
import argparse
import configparser
import csv
import glob
import json
import os
import re
import sys
import xml.etree.ElementTree as ET

import requests
import urllib3
import yaml

import cartridge_validate

urllib3.disable_warnings(urllib3.exceptions.InsecureRequestWarning)

HERE = os.path.dirname(os.path.abspath(__file__))
CHUNK = 18000  # /knowledge-convert 가 20000 에서 자르므로 여유를 두고 청크

# 구조화 소스에서 Q&A 열을 찾을 때 쓰는 후보 (대소문자 무시)
Q_KEYS = ("question", "q", "질문", "title", "제목")
A_KEYS = ("answer", "a", "답변", "content", "body", "본문", "내용")
AL_KEYS = ("aliases", "alias", "변형", "동의어")

# 추가 라이브러리가 필요한 포맷 → 설치 안내 (미설치 시 그 파일만 건너뜀)
_LAZY_HINT = {
    "pdf": "pip install pdfplumber",
    "docx": "pip install python-docx",
    "xlsx": "pip install openpyxl",
    "ocr": "pip install pytesseract pillow  (+ 시스템 tesseract-ocr)",
}


# ── 포맷별 텍스트/구조 추출 ──────────────────────────────────────────
def _pick(d, keys):
    """dict 키를 대소문자·후보군으로 관대하게 찾는다."""
    low = {str(k).strip().lower(): k for k in d.keys()}
    for cand in keys:
        if cand in low:
            return d[low[cand]]
    return None


def _rows_to_items(rows):
    """dict 행 목록에서 Q&A 열이 있으면 구조화 항목으로, 없으면 None(→텍스트 폴백)."""
    items = []
    for r in rows:
        q = _pick(r, Q_KEYS)
        a = _pick(r, A_KEYS)
        if not (q and str(q).strip() and a and str(a).strip()):
            return None  # 한 행이라도 Q&A 아님 → 구조화 아님, 텍스트로 처리
        al = _pick(r, AL_KEYS)
        aliases = ([s.strip() for s in re.split(r"[|;,]", str(al)) if s.strip()]
                   if al else [])
        items.append({"question": str(q).strip(), "answer": str(a).strip(),
                      "aliases": aliases})
    return items or None


def extract(path):
    """(kind, payload) 반환. kind='items' → 구조화 직접 매핑, kind='text' → LLM 변환 대상.

    추출 실패(라이브러리 미설치 등)는 (None, 사유) 로 돌려 호출부가 건너뛴다.
    """
    ext = os.path.splitext(path)[1].lower()
    try:
        if ext in (".md", ".markdown", ".txt"):
            with open(path, encoding="utf-8", errors="replace") as f:
                return "text", f.read()

        if ext == ".csv":
            with open(path, encoding="utf-8", errors="replace", newline="") as f:
                rows = list(csv.DictReader(f))
            items = _rows_to_items(rows) if rows else None
            if items:
                return "items", items
            return "text", "\n".join(" ".join(str(v) for v in r.values()) for r in rows)

        if ext == ".json":
            with open(path, encoding="utf-8", errors="replace") as f:
                data = json.load(f)
            if isinstance(data, list) and data and all(isinstance(x, dict) for x in data):
                items = _rows_to_items(data)
                if items:
                    return "items", items
            return "text", json.dumps(data, ensure_ascii=False, indent=2)

        if ext == ".xml":
            root = ET.parse(path).getroot()
            return "text", "\n".join(t.strip() for t in root.itertext() if t and t.strip())

        if ext == ".pdf":
            try:
                import pdfplumber
            except ImportError:
                return None, f"pdf 추출기 미설치 ({_LAZY_HINT['pdf']})"
            with pdfplumber.open(path) as pdf:
                return "text", "\n".join((p.extract_text() or "") for p in pdf.pages)

        if ext == ".docx":
            try:
                import docx
            except ImportError:
                return None, f"docx 추출기 미설치 ({_LAZY_HINT['docx']})"
            return "text", "\n".join(p.text for p in docx.Document(path).paragraphs)

        if ext == ".xlsx":
            try:
                import openpyxl
            except ImportError:
                return None, f"xlsx 추출기 미설치 ({_LAZY_HINT['xlsx']})"
            wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
            rows = []
            for ws in wb.worksheets:
                it = ws.iter_rows(values_only=True)
                header = next(it, None)
                if not header:
                    continue
                cols = [str(h).strip() if h is not None else f"col{i}"
                        for i, h in enumerate(header)]
                for row in it:
                    rows.append({cols[i]: row[i] for i in range(min(len(cols), len(row)))})
            items = _rows_to_items(rows) if rows else None
            if items:
                return "items", items
            return "text", "\n".join(" ".join(str(v) for v in r.values()) for r in rows)

        if ext in (".png", ".jpg", ".jpeg", ".tiff", ".tif", ".bmp"):
            try:
                import pytesseract
                from PIL import Image
            except ImportError:
                return None, f"ocr 미설치 ({_LAZY_HINT['ocr']})"
            return "text", pytesseract.image_to_string(Image.open(path), lang="kor+eng")

        return None, f"미지원 확장자 {ext}"
    except Exception as e:  # 손상 파일 등 — 그 파일만 건너뜀
        return None, f"{type(e).__name__}: {e}"


def _chunks(text):
    """긴 텍스트를 문단 경계 기준 ≤CHUNK 조각으로."""
    text = text.strip()
    if len(text) <= CHUNK:
        return [text] if text else []
    out, buf = [], ""
    for para in text.split("\n\n"):
        if len(buf) + len(para) + 2 > CHUNK and buf:
            out.append(buf)
            buf = ""
        buf = (buf + "\n\n" + para).strip() if buf else para
        while len(buf) > CHUNK:  # 문단 하나가 너무 크면 강제 절단
            out.append(buf[:CHUNK])
            buf = buf[CHUNK:]
    if buf.strip():
        out.append(buf)
    return out


# ── 로컬 LLM 변환 (기존 통로 재사용) ────────────────────────────────
def llm_convert(base_url, key, text, count):
    """비정형 텍스트 → Q&A 항목. /api/wizard/knowledge-convert 를 청크별로 호출."""
    items = []
    for ch in _chunks(text):
        r = requests.post(f"{base_url}/api/wizard/knowledge-convert",
                          headers={"Authorization": f"Bearer {key}"},
                          json={"text": ch, "count": count}, verify=False, timeout=300)
        r.raise_for_status()
        items.extend(r.json().get("items", []))
    return items


# ── 카트리지 쓰기 (위저드 cartridge-save 와 동일 포맷) ────────────────
def write_cartridge(name, items, system_prompt, description):
    base = os.path.join(HERE, "cartridges", name)
    os.makedirs(os.path.join(base, "prompts"))
    os.makedirs(os.path.join(base, "knowledge"))

    with open(os.path.join(base, "prompts", "system.txt"), "w", encoding="utf-8") as f:
        f.write(system_prompt.strip() + "\n")

    files = []
    for i, it in enumerate(items, 1):
        slug = re.sub(r"[^a-z0-9]+", "-", it["question"].lower())[:40].strip("-") or "item"
        fname = f"{i:03d}-{slug}.yaml"
        doc = {"type": "qna", "name": f"{name}-{i:03d}", "enabled": True,
               "question": it["question"], "answer": it["answer"],
               "aliases": it.get("aliases", [])}
        with open(os.path.join(base, "knowledge", fname), "w", encoding="utf-8") as f:
            yaml.safe_dump(doc, f, allow_unicode=True, sort_keys=False,
                           default_flow_style=False)
        files.append(fname)

    manifest = {
        "name": name, "version": "0.1.0", "description": description,
        "prompts": {"system": "prompts/system.txt"},
        "knowledge": {"dir": "knowledge/"},
        "model": {"recommended": "(설치 프리셋)"},
    }
    with open(os.path.join(base, "cartridge.yaml"), "w", encoding="utf-8") as f:
        yaml.safe_dump(manifest, f, allow_unicode=True, sort_keys=False)
    return base, files


def _default_base_url():
    cfg = configparser.ConfigParser()
    cfg.read(os.path.join(HERE, "config.ini"))
    port = cfg.get("server", "port", fallback="443").strip()
    return f"https://localhost:{port}"


def _default_key():
    try:
        with open(os.path.join(HERE, "api_keys", "default.key"), encoding="utf-8") as f:
            return f.read().strip()
    except OSError:
        return ""


def main():
    ap = argparse.ArgumentParser(description="문서 디렉터리 → 카트리지 지식 배치 인제스트")
    ap.add_argument("src_dir")
    ap.add_argument("name", help="카트리지 이름 (kebab-case)")
    ap.add_argument("--count", type=int, default=6, help="비정형 청크당 Q&A 개수")
    ap.add_argument("--base-url", default=None)
    ap.add_argument("--key", default=None)
    ap.add_argument("--system-prompt-file", default=None)
    ap.add_argument("--dry-run", action="store_true")
    args = ap.parse_args()

    if not re.fullmatch(r"[a-z0-9][a-z0-9-]{1,40}", args.name):
        sys.exit("✗ 카트리지 이름은 kebab-case (소문자·숫자·하이픈, 2~41자)")
    dst = os.path.join(HERE, "cartridges", args.name)
    if not args.dry_run and os.path.exists(dst):
        sys.exit(f"✗ 이미 존재: {dst} (다른 이름을 쓰거나 먼저 지우세요)")
    if not os.path.isdir(args.src_dir):
        sys.exit(f"✗ 소스 디렉터리 없음: {args.src_dir}")

    base_url = (args.base_url or _default_base_url()).rstrip("/")
    key = args.key or _default_key()
    if not key:
        sys.exit("✗ Bearer 키를 찾지 못함 (api_keys/default.key 또는 --key)")

    paths = sorted(p for p in glob.glob(os.path.join(args.src_dir, "**", "*"), recursive=True)
                   if os.path.isfile(p))
    print(f"📂 {len(paths)}개 파일 스캔 · 콘솔 {base_url}")

    items, skipped, structured_n, llm_n = [], [], 0, 0
    for p in paths:
        kind, payload = extract(p)
        rel = os.path.relpath(p, args.src_dir)
        if kind is None:
            skipped.append((rel, payload))
            print(f"  – 건너뜀 {rel} — {payload}")
            continue
        if kind == "items":
            items.extend(payload)
            structured_n += len(payload)
            print(f"  ✓ {rel} — 구조화 {len(payload)}건 (직접 매핑)")
        else:
            try:
                got = llm_convert(base_url, key, payload, args.count)
            except Exception as e:
                skipped.append((rel, f"변환 실패 {type(e).__name__}: {e}"))
                print(f"  ✗ {rel} — 변환 실패: {e}")
                continue
            items.extend(got)
            llm_n += len(got)
            print(f"  ✓ {rel} — LLM 변환 {len(got)}건 (초안·검수 권장)")

    print(f"\n합계 {len(items)}건 (구조화 {structured_n} + LLM {llm_n}) · 건너뜀 {len(skipped)}")
    if not items:
        sys.exit("✗ 추출된 Q&A 항목이 없습니다.")
    if args.dry_run:
        print("(--dry-run: 파일을 쓰지 않았습니다)")
        return

    sp = "You are a helpful assistant for this domain."
    if args.system_prompt_file:
        with open(args.system_prompt_file, encoding="utf-8") as f:
            sp = f.read()
    base, files = write_cartridge(args.name, items, sp,
                                  f"{args.name} — ingest.py 로 생성 ({len(items)}건)")
    print(f"📦 저장: {base} (지식 {len(files)}개)")

    rep = cartridge_validate.validate_cartridge(base)
    print(f"🔎 validate — 오류 {rep.n_errors} · 경고 {rep.n_warnings}")
    for c in rep.checks:
        if c["status"] != cartridge_validate.OK:
            print(f"    {c['status']} {c['label']}: {c['detail']}")

    print("\n다음:")
    print(f"  1) 검수 — cartridges/{args.name}/knowledge/*.yaml (특히 LLM 변환분)")
    print(f"  2) 장착 — ./aibotctl cartridge mount {args.name}")
    print("  ⚠️ 이 카트리지는 prompts.system 만 있는 단일형이라 intent 모드에서 페르소나는")
    print("     기본값입니다. 지식(RAG)은 항상 물립니다. 페르소나가 필요하면 prompts.qna 를")
    print("     추가하거나 위저드 프롬프트 초안(/api/wizard/prompt-draft)을 쓰세요.")


if __name__ == "__main__":
    main()
